from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def write_table_excel(path: Path | str, title: str, headers: list[str], rows: list[list[str]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31] or "Report"

    header_fill = PatternFill(start_color="1E2536", end_color="1E2536", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    ws.append(headers)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font

    for row in rows:
        ws.append(row)

    for column_cells in ws.columns:
        length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
        ws.column_dimensions[column_cells[0].column_letter].width = max(10, length + 2)

    wb.save(str(path))
